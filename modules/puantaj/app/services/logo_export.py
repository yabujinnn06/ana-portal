"""Logo aktarim dosyalari: Bordro Plus puantaj aktarimi + Tiger/GO muhasebe mahsup fisi.

Iki ayri is:
- Puantaj aktarimi: calisan basina donem ozeti (SGK gun, FM saatleri, eksik gun). Logo
  Bordro Plus / IK bu veriyle KENDI bordrosunu hesaplar.
- Muhasebe mahsup fisi: bu app'in hesapladigi bordroyu Logo ERP'ye muhasebe fisi (borc/alacak)
  olarak yazar. Hesap kodlari CompanySettings'ten gelir, yoksa TR varsayilan kullanilir.
"""

from __future__ import annotations

from calendar import monthrange
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.models import SgkStatus
from app.services.exports import (
    THIN_BORDER,
    ZEBRA_FILL,
    _apply_print_layout,
    _apply_workbook_branding,
    _append_spacer_row,
    _auto_width,
    _style_header,
    _style_metadata_rows,
    _to_excel_datetime,
)

MONEY_PLACES = Decimal("0.01")
HOUR_PLACES = Decimal("0.01")
MONEY_FORMAT = "#,##0.00"

# Logo muhasebe mahsup fisi varsayilan hesap kodlari (TR tek duzen hesap plani).
DEFAULT_LOGO_ACCOUNTS = {
    "ucret": "770",
    "sgk_isveren": "770",
    "net_odenecek": "335",
    "odenecek_vergi": "360",
    "odenecek_sgk": "361",
    "personel_kesinti": "335",
}


def _money(value: object) -> Decimal:
    return Decimal(value or 0).quantize(MONEY_PLACES, rounding=ROUND_HALF_UP)


def _hours(minutes: object) -> Decimal:
    return (Decimal(max(0, int(minutes or 0))) / Decimal("60")).quantize(
        HOUR_PLACES, rounding=ROUND_HALF_UP
    )


def _account(settings: object, key: str) -> str:
    value = getattr(settings, f"logo_hesap_{key}", None) if settings is not None else None
    return (str(value).strip() if value else "") or DEFAULT_LOGO_ACCOUNTS[key]


def _period_end(year: int, month: int) -> date:
    return date(year, month, monthrange(year, month)[1])


# --------------------------------------------------------------------------- #
# Logo Bordro Plus / IK - PUANTAJ AKTARIMI
# --------------------------------------------------------------------------- #

PUANTAJ_HEADERS = [
    "Sira",
    "Sicil No",
    "TC Kimlik No",
    "Ad Soyad",
    "Departman",
    "SGK Statu",
    "Kanun No",
    "SGK Prim Gun",
    "Eksik Gun (SGK)",
    "Normal Calisma (saat)",
    "FM %50 (saat)",
    "Resmi/Bayram FM %100 (saat)",
    "Hafta Tatili FM %100 (saat)",
    "Ucretsiz Izin (saat)",
    "Eksik Calisma (saat)",
    "Brut Hakedis (TL)",
]


def build_logo_puantaj_xlsx_bytes(run, items, profiles_by_employee, *, generated_at: datetime) -> bytes:
    """Logo Bordro Plus puantaj aktarimi: calisan basina donem ozeti (bir satir)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Logo Puantaj"

    metadata_rows = [
        ("Donem", f"{run.year}-{run.month:02d}"),
        ("Durum", "ONAYLI" if run.status.value == "APPROVED" else "TASLAK"),
        ("Olusturma", _to_excel_datetime(generated_at)),
        ("Kayit Sayisi", len(items)),
        ("Aciklama", "Logo Bordro Plus puantaj aktarimi"),
    ]
    for label, value in metadata_rows:
        ws.append([label, value])
    _style_metadata_rows(ws, start_row=1, end_row=len(metadata_rows))
    ws["B3"].number_format = "yyyy-mm-dd hh:mm"

    _append_spacer_row(ws)
    header_row = ws.max_row + 1
    ws.append(PUANTAJ_HEADERS)
    _style_header(ws, header_row)
    ws.freeze_panes = f"A{header_row + 1}"

    for index, item in enumerate(items, start=1):
        profile = profiles_by_employee.get(item.employee_id)
        sicil = (getattr(profile, "sgk_sicil_no", None) if profile else None) or item.employee_id
        tc = (getattr(profile, "tc_kimlik_no", None) or "-") if profile else "-"
        sgk_days = max(0, int(item.sgk_days))
        eksik_gun = max(0, 30 - sgk_days)
        sgk_label = "Emekli (SGDP)" if item.sgk_status == SgkStatus.EMEKLI else "Normal (4a)"
        ws.append(
            [
                index,
                sicil,
                tc,
                item.employee_name,
                item.department_name or "-",
                sgk_label,
                item.kanun_no or "-",
                sgk_days,
                eksik_gun,
                _hours(item.worked_minutes),
                _hours(item.fm1_minutes),
                _hours(item.fm2_minutes),
                _hours(item.fm3_minutes),
                _hours(item.unpaid_leave_minutes),
                _hours(item.missing_minutes),
                _money(Decimal(item.gross_total) + Decimal(item.additional_earnings)),
            ]
        )

    data_end_row = ws.max_row
    for row_idx in range(header_row + 1, data_end_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = THIN_BORDER
            if row_idx % 2 == 0:
                cell.fill = ZEBRA_FILL
        ws.cell(row=row_idx, column=16).number_format = MONEY_FORMAT

    _auto_width(ws)
    _apply_print_layout(ws, header_row=header_row)
    _apply_workbook_branding(wb)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


# --------------------------------------------------------------------------- #
# Logo Tiger / GO - MUHASEBE MAHSUP FISI
# --------------------------------------------------------------------------- #

MAHSUP_HEADERS = [
    "Fis Tarihi",
    "Fis No",
    "Satir",
    "Hesap Kodu",
    "Hesap Adi",
    "Aciklama",
    "Borc",
    "Alacak",
]


def _mahsup_lines(items, settings, *, description: str) -> list[dict]:
    """Toplu mahsup satirlari. Borc toplam == Alacak toplam (item snapshot ile dengeli)."""

    def total(attr: str) -> Decimal:
        return sum((Decimal(getattr(i, attr)) for i in items), Decimal("0"))

    ucret = total("gross_total") + total("additional_earnings")
    sgk_isveren = total("sgk_employer") + total("unemployment_employer")
    net = total("net_total")
    vergi = total("income_tax_payable") + total("stamp_tax_payable")
    sgk_kesinti = (
        total("sgk_employee")
        + total("unemployment_employee")
        + total("sgk_employer")
        + total("unemployment_employer")
    )
    personel_kesinti = total("additional_deductions")

    rows = [
        ("ucret", "Ucret Giderleri", "Ucret Tahakkuku", _money(ucret), Decimal("0")),
        ("sgk_isveren", "SGK Isveren Payi Gideri", "SGK + Issizlik Isveren Payi", _money(sgk_isveren), Decimal("0")),
        ("net_odenecek", "Personele Borclar", "Net Ucret Odenecek", Decimal("0"), _money(net)),
        ("odenecek_vergi", "Odenecek Vergi ve Fonlar", "Gelir Vergisi + Damga", Decimal("0"), _money(vergi)),
        ("odenecek_sgk", "Odenecek SGK Kesintileri", "SGK Isci + Isveren", Decimal("0"), _money(sgk_kesinti)),
        ("personel_kesinti", "Personel Kesintileri", "Avans / Icra Kesintisi", Decimal("0"), _money(personel_kesinti)),
    ]

    lines: list[dict] = []
    for key, name, line_desc, borc, alacak in rows:
        if borc == 0 and alacak == 0:
            continue
        lines.append(
            {
                "hesap_kodu": _account(settings, key),
                "hesap_adi": name,
                "aciklama": f"{description} - {line_desc}",
                "borc": borc,
                "alacak": alacak,
            }
        )
    return lines


def build_logo_mahsup_xlsx_bytes(run, items, settings, *, generated_at: datetime) -> bytes:
    """Logo ERP muhasebe mahsup fisi (Excel): donem sonu tarihli toplu bordro fisi."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Mahsup Fisi"

    fis_tarihi = _period_end(run.year, run.month)
    fis_no = f"BORDRO-{run.year}{run.month:02d}"
    description = f"{run.year}-{run.month:02d} Donem Bordro Tahakkuku"
    lines = _mahsup_lines(items, settings, description=description)
    toplam_borc = sum((line["borc"] for line in lines), Decimal("0"))
    toplam_alacak = sum((line["alacak"] for line in lines), Decimal("0"))

    metadata_rows = [
        ("Firma", (getattr(settings, "firma_unvan", None) or "-") if settings else "-"),
        ("Donem", f"{run.year}-{run.month:02d}"),
        ("Fis Turu", "Mahsup"),
        ("Fis Tarihi", fis_tarihi),
        ("Fis No", fis_no),
        ("Durum", "ONAYLI" if run.status.value == "APPROVED" else "TASLAK"),
        ("Olusturma", _to_excel_datetime(generated_at)),
        ("Denge", "DENGELI" if toplam_borc == toplam_alacak else "UYUMSUZ"),
    ]
    for label, value in metadata_rows:
        ws.append([label, value])
    _style_metadata_rows(ws, start_row=1, end_row=len(metadata_rows))
    ws["B4"].number_format = "yyyy-mm-dd"
    ws["B7"].number_format = "yyyy-mm-dd hh:mm"

    _append_spacer_row(ws)
    header_row = ws.max_row + 1
    ws.append(MAHSUP_HEADERS)
    _style_header(ws, header_row)
    ws.freeze_panes = f"A{header_row + 1}"

    for line_idx, line in enumerate(lines, start=1):
        ws.append(
            [
                fis_tarihi,
                fis_no,
                line_idx,
                line["hesap_kodu"],
                line["hesap_adi"],
                line["aciklama"],
                line["borc"] if line["borc"] != 0 else None,
                line["alacak"] if line["alacak"] != 0 else None,
            ]
        )

    data_end_row = ws.max_row
    for row_idx in range(header_row + 1, data_end_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = THIN_BORDER
            if row_idx % 2 == 0:
                cell.fill = ZEBRA_FILL
        ws.cell(row=row_idx, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=row_idx, column=7).number_format = MONEY_FORMAT
        ws.cell(row=row_idx, column=8).number_format = MONEY_FORMAT

    if lines:
        total_row = data_end_row + 1
        ws.append([None, None, None, None, None, "TOPLAM", toplam_borc, toplam_alacak])
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.font = Font(bold=True)
            cell.border = THIN_BORDER
        ws.cell(row=total_row, column=7).number_format = MONEY_FORMAT
        ws.cell(row=total_row, column=8).number_format = MONEY_FORMAT

    _auto_width(ws)
    _apply_print_layout(ws, header_row=header_row)
    _apply_workbook_branding(wb)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
