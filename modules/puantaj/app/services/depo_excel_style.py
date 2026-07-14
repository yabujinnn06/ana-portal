"""Depo/sayim Excel export stil yardimcilari.

Ported verbatim from the standalone depojin app (backend/app/excel_style.py) as part of
the depo entegrasyonu (see CLAUDE.md). No behavior changes.
"""

from __future__ import annotations

from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage, ImageDraw, ImageFont

DEEP = "FF0F2A44"
ACCENT = "FFBF6F34"
ACCENT_BG = "FFFFEAD6"
GOOD = "FF5FBE7A"
GOOD_BG = "FFE6F4DC"
WARN_BG = "FFFCE5CB"
BAD = "FFDC5A5A"
BAD_BG = "FFFCE0E0"
CREAM = "FFF4F0E8"
EDGE = "FFC6BDAC"
GRUP_BASLIK_BG = "FFDDEBF7"


def border(thin: bool = True) -> Border:
    s = Side(style="thin", color=EDGE) if thin else Side(style="medium", color=EDGE)
    return Border(left=s, right=s, top=s, bottom=s)


_FONT_YOLLARI = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf",
    "/Library/Fonts/Arial Bold.ttf",
    "C:/Windows/Fonts/arialbd.ttf",
]


def _font(boyut: int):
    for p in _FONT_YOLLARI:
        try:
            return ImageFont.truetype(p, boyut)
        except Exception:
            continue
    return ImageFont.load_default()


_FILIGRAN_CACHE: bytes | None = None


def _filigran_png() -> bytes:
    global _FILIGRAN_CACHE
    if _FILIGRAN_CACHE is not None:
        return _FILIGRAN_CACHE
    W, H = 900, 280
    img = PILImage.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    f1 = _font(140)
    f2 = _font(34)
    txt1 = "YABUJIN"
    txt2 = "rainwater sayim sistemi"
    try:
        b1 = d.textbbox((0, 0), txt1, font=f1)
        b2 = d.textbbox((0, 0), txt2, font=f2)
        w1, h1 = b1[2] - b1[0], b1[3] - b1[1]
        w2, h2 = b2[2] - b2[0], b2[3] - b2[1]
    except Exception:
        w1 = h1 = w2 = h2 = 0
    d.text(((W - w1) / 2, (H - h1) / 2 - 18), txt1, fill=(191, 111, 52, 28), font=f1)
    d.text(((W - w2) / 2, (H - h1) / 2 + h1 - 8), txt2, fill=(15, 42, 68, 22), font=f2)
    buf = BytesIO()
    img.save(buf, format="PNG")
    _FILIGRAN_CACHE = buf.getvalue()
    return _FILIGRAN_CACHE


def filigran_ekle(ws, satir_aralik: int = 22) -> None:
    son_sat = ws.max_row or 1
    son_kol = ws.max_column or 8
    if son_sat < 1:
        return
    konum_satirlari = list(range(8, son_sat + satir_aralik, satir_aralik))
    if not konum_satirlari:
        konum_satirlari = [8]
    for sat in konum_satirlari:
        png = BytesIO(_filigran_png())
        xl = XLImage(png)
        xl.width = 520
        xl.height = 160
        kol_anchor = max(0, (son_kol // 2) - 3)
        xl.anchor = f"{get_column_letter(kol_anchor + 1)}{sat}"
        ws.add_image(xl)


def baslik_yaz(ws, satir, kolonlar, baslik_renk=DEEP, yazi_renk="FFFFFFFF"):
    for col_idx, (etiket, _w) in enumerate(kolonlar, start=1):
        c = ws.cell(row=satir, column=col_idx, value=etiket)
        c.font = Font(bold=True, color=yazi_renk, size=11)
        c.fill = PatternFill("solid", fgColor=baslik_renk)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border()


def kolon_genislikleri(ws, kolonlar):
    for i, (_e, w) in enumerate(kolonlar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def baslik_satiri_yaz(ws, satir, kolon_sayisi, metin, renk=DEEP, boyut=14, yukseklik=28):
    ws.merge_cells(start_row=satir, start_column=1, end_row=satir, end_column=kolon_sayisi)
    c = ws.cell(row=satir, column=1, value=metin)
    c.font = Font(bold=True, color="FFFFFFFF", size=boyut)
    c.fill = PatternFill("solid", fgColor=renk)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[satir].height = yukseklik


def grup_baslik_hazirla(ws, satir, kolon_sayisi):
    """Stok bazli gorunum icin acik mavi grup basligi satiri."""
    for col in range(1, kolon_sayisi + 1):
        c = ws.cell(row=satir, column=col)
        c.fill = PatternFill("solid", fgColor=GRUP_BASLIK_BG)
        c.font = Font(bold=True, size=11)
        c.border = border()


def gruplamayi_etkinlestir(ws):
    """+/- ile acilip kapanan stok bazli gorunum: detay satirlari varsayilan gizli."""
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_format.outlineLevelRow = 1


def detay_satiri_grupla(ws, satir):
    rd = ws.row_dimensions[satir]
    rd.outlineLevel = 1
    rd.hidden = True
