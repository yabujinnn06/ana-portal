"""Sayim oturumu icin log filtreleme/istatistik/eksik seri raporlari.

Ported from the standalone depojin app (backend/app/routers/rapor.py). Original prefix was
``/api/sayim`` (shares the prefix with sayim.py); per the depo entegrasyonu mapping it stays
``/sayim`` here, landing at ``/api/depo/sayim/...`` alongside sayim.py's own routes (no path
overlap between the two modules).
"""

from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from collections import defaultdict
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select, func, and_, or_, case
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import SayimOturumu, DepoSayimStok, DepoSeri, DepoTaramaLog, DepoUser
from app.depo_schemas import (
    LogOut, LogSayfaOut, IstatistikOut, DurumSayim,
    KullaniciIstatistik, DakikaSayim, EksikGrupOut,
)
from app.services.depo_auth import current_depo_user
from app.services.depo_excel_style import (
    DEEP, GOOD_BG, WARN_BG, BAD_BG,
    border as _border, baslik_yaz as _baslik_yaz, kolon_genislikleri as _kolon_genislikleri,
    baslik_satiri_yaz as _baslik_satiri_yaz, filigran_ekle as _filigran_ekle,
    grup_baslik_hazirla, gruplamayi_etkinlestir, detay_satiri_grupla,
)

router = APIRouter(prefix="/sayim", tags=["depo-rapor"])


def _kontrol(db: Session, oturum_id: int) -> SayimOturumu:
    o = db.get(SayimOturumu, oturum_id)
    if not o:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Oturum yok")
    return o


@router.get("/{oturum_id}/log-filtre", response_model=LogSayfaOut)
def log_filtre(
    oturum_id: int,
    durum: str | None = Query(None),
    kullanici_id: int | None = Query(None),
    q: str | None = Query(None),
    baslangic: datetime | None = Query(None),
    bitis: datetime | None = Query(None),
    limit: int = Query(100, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    db: Session = Depends(get_db),
    _: DepoUser = Depends(current_depo_user),
):
    _kontrol(db, oturum_id)
    kosul = [DepoTaramaLog.oturum_id == oturum_id]
    if durum:
        kosul.append(DepoTaramaLog.durum == durum)
    if kullanici_id is not None:
        kosul.append(DepoTaramaLog.kullanici_id == kullanici_id)
    if baslangic:
        kosul.append(DepoTaramaLog.zaman >= baslangic)
    if bitis:
        kosul.append(DepoTaramaLog.zaman <= bitis)
    if q:
        like = f"%{q.strip()}%"
        kosul.append(or_(
            DepoTaramaLog.seri_giris.ilike(like),
            DepoTaramaLog.stok_kodu.ilike(like),
            DepoTaramaLog.urun_adi.ilike(like),
            DepoTaramaLog.aciklama.ilike(like),
        ))

    toplam = db.scalar(select(func.count(DepoTaramaLog.id)).where(and_(*kosul))) or 0
    rows = db.execute(
        select(DepoTaramaLog, DepoUser.ad)
        .outerjoin(DepoUser, DepoUser.id == DepoTaramaLog.kullanici_id)
        .where(and_(*kosul))
        .order_by(DepoTaramaLog.zaman.desc())
        .offset(offset).limit(limit)
    ).all()

    items: list[LogOut] = []
    for log, ad in rows:
        it = LogOut.model_validate(log)
        it.kullanici_ad = ad
        items.append(it)
    return LogSayfaOut(toplam=toplam, items=items)


@router.get("/{oturum_id}/istatistik", response_model=IstatistikOut)
def istatistik(oturum_id: int, db: Session = Depends(get_db), _: DepoUser = Depends(current_depo_user)):
    _kontrol(db, oturum_id)

    durum_rows = db.execute(
        select(DepoTaramaLog.durum, func.count(DepoTaramaLog.id))
        .where(DepoTaramaLog.oturum_id == oturum_id)
        .group_by(DepoTaramaLog.durum)
        .order_by(func.count(DepoTaramaLog.id).desc())
    ).all()
    durum_dagilimi = [DurumSayim(durum=d, sayi=n) for d, n in durum_rows]

    kullanici_rows = db.execute(
        select(
            DepoTaramaLog.kullanici_id,
            DepoUser.ad,
            func.sum(case((DepoTaramaLog.durum == "basarili", 1), else_=0)),
            func.sum(case((DepoTaramaLog.durum == "mukerrer", 1), else_=0)),
            func.sum(case((DepoTaramaLog.durum == "bulunamadi", 1), else_=0)),
            func.sum(case((DepoTaramaLog.durum == "cakisma", 1), else_=0)),
            func.count(DepoTaramaLog.id),
            func.max(DepoTaramaLog.zaman),
        )
        .outerjoin(DepoUser, DepoUser.id == DepoTaramaLog.kullanici_id)
        .where(DepoTaramaLog.oturum_id == oturum_id)
        .group_by(DepoTaramaLog.kullanici_id, DepoUser.ad)
        .order_by(func.count(DepoTaramaLog.id).desc())
    ).all()
    kullanici_basina = [
        KullaniciIstatistik(
            kullanici_id=row[0],
            ad=row[1] or "Bilinmeyen",
            basarili=int(row[2] or 0),
            mukerrer=int(row[3] or 0),
            bulunamadi=int(row[4] or 0),
            cakisma=int(row[5] or 0),
            toplam_tarama=int(row[6] or 0),
            son_tarama=row[7],
        )
        for row in kullanici_rows
    ]

    ilk = db.scalar(select(func.min(DepoTaramaLog.zaman)).where(DepoTaramaLog.oturum_id == oturum_id))
    son = db.scalar(select(func.max(DepoTaramaLog.zaman)).where(DepoTaramaLog.oturum_id == oturum_id))

    dakika_serisi: list[DakikaSayim] = []
    if ilk and son:
        kova_basarili: dict[datetime, int] = defaultdict(int)
        kova_diger: dict[datetime, int] = defaultdict(int)
        loglar = db.execute(
            select(DepoTaramaLog.zaman, DepoTaramaLog.durum)
            .where(DepoTaramaLog.oturum_id == oturum_id)
        ).all()
        for z, d in loglar:
            kova = z.replace(second=0, microsecond=0)
            if d == "basarili":
                kova_basarili[kova] += 1
            else:
                kova_diger[kova] += 1
        anahtar = sorted(set(list(kova_basarili.keys()) + list(kova_diger.keys())))
        dakika_serisi = [
            DakikaSayim(zaman=z, basarili=kova_basarili[z], diger=kova_diger[z])
            for z in anahtar
        ]

    basarili_top = db.scalar(
        select(func.count(DepoTaramaLog.id))
        .where(and_(DepoTaramaLog.oturum_id == oturum_id, DepoTaramaLog.durum == "basarili"))
    ) or 0
    sure_dk = 0.0
    if ilk and son and basarili_top > 0:
        delta = (son - ilk).total_seconds() / 60.0
        sure_dk = (basarili_top / delta) if delta > 0 else float(basarili_top)
    tarama_dakika_dk = round(sure_dk, 2)

    return IstatistikOut(
        durum_dagilimi=durum_dagilimi,
        kullanici_basina=kullanici_basina,
        dakika_serisi=dakika_serisi,
        ilk_tarama=ilk,
        son_tarama=son,
        tarama_dakika_dk=tarama_dakika_dk,
    )


@router.get("/{oturum_id}/eksik", response_model=list[EksikGrupOut])
def eksik_seriler(
    oturum_id: int,
    limit_seri: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
    _: DepoUser = Depends(current_depo_user),
):
    _kontrol(db, oturum_id)
    stoklar = db.execute(
        select(DepoSayimStok).where(DepoSayimStok.oturum_id == oturum_id).order_by(DepoSayimStok.stok_kodu)
    ).scalars().all()

    sonuc: list[EksikGrupOut] = []
    for s in stoklar:
        toplam = db.scalar(select(func.count(DepoSeri.id)).where(DepoSeri.stok_id == s.id)) or 0
        sayilan = db.scalar(
            select(func.count(DepoSeri.id)).where(and_(DepoSeri.stok_id == s.id, DepoSeri.sayildi == True))
        ) or 0
        eksik = toplam - sayilan
        if eksik <= 0:
            continue
        seriler = db.execute(
            select(DepoSeri.seri_no)
            .where(and_(DepoSeri.stok_id == s.id, DepoSeri.sayildi == False))
            .order_by(DepoSeri.seri_no)
            .limit(limit_seri)
        ).scalars().all()
        sonuc.append(EksikGrupOut(
            stok_id=s.id, stok_kodu=s.stok_kodu, urun_adi=s.urun_adi,
            toplam=toplam, sayilan=sayilan, eksik=eksik,
            portal_sayim=s.portal_sayim, seriler=list(seriler),
        ))
    return sonuc


_DURUM_RENK = {
    "basarili": GOOD_BG, "mukerrer": WARN_BG, "bulunamadi": BAD_BG, "cakisma": WARN_BG,
}


@router.get("/{oturum_id}/log/excel")
def log_excel(oturum_id: int, db: Session = Depends(get_db), _: DepoUser = Depends(current_depo_user)):
    oturum = _kontrol(db, oturum_id)

    rows = db.execute(
        select(DepoTaramaLog, DepoUser.ad)
        .outerjoin(DepoUser, DepoUser.id == DepoTaramaLog.kullanici_id)
        .where(DepoTaramaLog.oturum_id == oturum_id)
        .order_by(DepoTaramaLog.zaman.asc())
    ).all()

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Kayitlar Gorunum"
    _gorunum_sayfasi(ws1, oturum, rows)
    ws2 = wb.create_sheet("Tum Kayitlar")
    _duz_liste_sayfasi(ws2, oturum, rows)
    for s in wb.worksheets:
        s.sheet_view.showGridLines = False
        _filigran_ekle(s)
    wb.properties.creator = "Yabujin · Rainwater"
    wb.properties.company = "Yabujin"
    wb.properties.title = f"Rainwater Kayit Raporu — {oturum.ad}"

    buf = BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"kayit_raporu_{oturum.ad.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


def _gorunum_sayfasi(ws, oturum: SayimOturumu, rows) -> None:
    """DepoSayim.xlsm > KayitlarGorunum ile ayni bicimde: stok bazli, +/- ile acilan gorunum."""
    kolonlar = [("Seri Giris", 24), ("Zaman", 19), ("Kullanici", 16), ("Durum", 12)]

    _baslik_satiri_yaz(
        ws, 1, len(kolonlar),
        f"KAYITLAR — STOK BAZLI GORUNUM — {oturum.ad}",
    )
    ws.row_dimensions[2].height = 6
    gruplamayi_etkinlestir(ws)
    _kolon_genislikleri(ws, kolonlar)

    gruplar: dict[str, list] = defaultdict(list)
    for log, ad in rows:
        gruplar[log.stok_kodu or "(Stoksuz)"].append((log, ad))

    satir = 3
    for stok_kodu in sorted(gruplar.keys()):
        kayitlar = gruplar[stok_kodu]
        urun_adi = next((log.urun_adi for log, _ in kayitlar if log.urun_adi), "") or ""

        grup_baslik_hazirla(ws, satir, len(kolonlar))
        ws.cell(row=satir, column=1, value=stok_kodu)
        ws.cell(row=satir, column=2, value=f"Adet: {len(kayitlar)}")
        ws.cell(row=satir, column=3, value=urun_adi)
        satir += 1

        for log, ad in kayitlar:
            ws.cell(row=satir, column=1, value=log.seri_giris)
            ws.cell(row=satir, column=2, value=log.zaman.strftime("%d.%m.%Y %H:%M:%S"))
            ws.cell(row=satir, column=3, value=ad or "")
            dc = ws.cell(row=satir, column=4, value=log.durum)
            bg = _DURUM_RENK.get(log.durum)
            if bg:
                dc.fill = PatternFill("solid", fgColor=bg)
                dc.font = Font(bold=True, size=10)
            for col in range(1, len(kolonlar) + 1):
                ws.cell(row=satir, column=col).border = _border()
            detay_satiri_grupla(ws, satir)
            satir += 1


def _duz_liste_sayfasi(ws, oturum: SayimOturumu, rows) -> None:
    kolonlar = [
        ("Zaman", 19), ("Kullanici", 16), ("Seri Giris", 22),
        ("Durum", 12), ("Stok Kodu", 14), ("Urun Adi", 32), ("Aciklama", 30),
    ]
    _baslik_yaz(ws, 1, kolonlar)
    _kolon_genislikleri(ws, kolonlar)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    sat = 2
    for log, ad in rows:
        ws.cell(row=sat, column=1, value=log.zaman.strftime("%d.%m.%Y %H:%M:%S"))
        ws.cell(row=sat, column=2, value=ad or "")
        ws.cell(row=sat, column=3, value=log.seri_giris)
        dc = ws.cell(row=sat, column=4, value=log.durum)
        bg = _DURUM_RENK.get(log.durum)
        if bg:
            dc.fill = PatternFill("solid", fgColor=bg)
            dc.font = Font(bold=True)
        ws.cell(row=sat, column=5, value=log.stok_kodu or "")
        ws.cell(row=sat, column=6, value=log.urun_adi or "")
        ws.cell(row=sat, column=7, value=log.aciklama or "")
        for col in range(1, len(kolonlar) + 1):
            ws.cell(row=sat, column=col).border = _border()
        sat += 1
