"""Sayim oturumu icin bicimlendirilmis Excel raporu.

Ported from the standalone depojin app (backend/app/routers/export.py). Original prefix
was ``/api/export``; per the depo entegrasyonu mapping it becomes ``/export`` here, landing
at ``/api/depo/export/...``.
"""

from __future__ import annotations

from io import BytesIO
from typing import Sequence
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import select, func, and_
from sqlalchemy.orm import Session, selectinload

from app.db import get_db
from app.models import SayimOturumu, DepoSayimStok, DepoSeri, DepoUser, Employee
from app.services.depo_auth import current_depo_user
from app.services.depo_excel_style import (
    DEEP, ACCENT, ACCENT_BG, GOOD, GOOD_BG, WARN_BG, BAD, BAD_BG,
    GRUP_BASLIK_BG,
    border as _border, baslik_yaz as _baslik_yaz, kolon_genislikleri as _kolon_genislikleri,
    filigran_ekle as _filigran_ekle, gruplamayi_etkinlestir, detay_satiri_grupla,
)

router = APIRouter(prefix="/export", tags=["depo-export"])


@router.get("/{oturum_id}/excel")
def export_excel(oturum_id: int, db: Session = Depends(get_db), _: DepoUser = Depends(current_depo_user)):
    oturum = db.get(SayimOturumu, oturum_id)
    if not oturum:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Oturum yok")

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Sayim"
    _ozet_sayfasi(db, wb, oturum)
    _sayim_sayfasi(db, ws1, oturum)
    _log_sayfasi(db, wb, oturum)
    for s in wb.worksheets:
        s.sheet_view.showGridLines = False
        _filigran_ekle(s)
    wb.properties.creator = "Yabujin · Rainwater"
    wb.properties.company = "Yabujin"
    wb.properties.title = f"Rainwater Sayim — {oturum.ad}"
    wb.properties.subject = "Sayim raporu"

    buf = BytesIO()
    wb.save(buf); buf.seek(0)
    fname = f"sayim_{oturum.ad.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


def _zimmetli(seri: DepoSeri) -> bool:
    return bool(seri.zimmet_employee_id or seri.zimmet_kullanici_id)


def _zimmet_ad(seri: DepoSeri, adlar: dict[int, str], employee_adlar: dict[int, str]) -> str | None:
    if seri.zimmet_employee_id:
        return employee_adlar.get(seri.zimmet_employee_id)
    if seri.zimmet_kullanici_id:
        return adlar.get(seri.zimmet_kullanici_id)
    return None


def _employee_adlari(db: Session, stoklar: Sequence[DepoSayimStok]) -> dict[int, str]:
    ids = {s.zimmet_employee_id for stok in stoklar for s in stok.seriler if s.zimmet_employee_id}
    if not ids:
        return {}
    return {e.id: e.full_name for e in db.execute(select(Employee).where(Employee.id.in_(ids))).scalars()}


def _sayim_sayfasi(db: Session, ws, oturum: SayimOturumu) -> None:
    kolonlar = [
        ("Stok Kodu", 14), ("Urun Adi", 32), ("Seri No", 22),
        ("Portal Sayi", 12), ("Sayilan", 12), ("Fark", 9),
        ("Sayim Tarihi", 19), ("Sayan", 16),
        ("Durum", 12), ("Sonradan Eklendi", 11), ("Cikis Notu", 22),
        ("Zimmet Hedefi", 18), ("Zimmet Notu", 22), ("Notlar", 24),
    ]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(kolonlar))
    c0 = ws.cell(row=1, column=1, value=f"SAYIM RAPORU — {oturum.ad}")
    c0.font = Font(bold=True, color="FFFFFFFF", size=14)
    c0.fill = PatternFill("solid", fgColor=DEEP)
    c0.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    meta = f"Lokasyon: {oturum.lokasyon or '-'}   Durum: {oturum.durum}   Baslangic: {oturum.baslangic.strftime('%d.%m.%Y %H:%M')}"
    if oturum.bitis:
        meta += f"   Bitis: {oturum.bitis.strftime('%d.%m.%Y %H:%M')}"
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(kolonlar))
    c1 = ws.cell(row=2, column=1, value=meta)
    c1.font = Font(italic=True, color="FF555555", size=10)
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    _baslik_yaz(ws, 3, kolonlar)
    _kolon_genislikleri(ws, kolonlar)
    ws.row_dimensions[3].height = 32
    ws.freeze_panes = "D4"
    gruplamayi_etkinlestir(ws)

    satir = 4
    stoklar = db.execute(
        select(DepoSayimStok)
        .where(DepoSayimStok.oturum_id == oturum.id)
        .order_by(DepoSayimStok.stok_kodu)
        .options(selectinload(DepoSayimStok.seriler))
    ).scalars().all()
    adlar = {u.id: u.ad for u in db.execute(select(DepoUser)).scalars()}
    employee_adlar = _employee_adlari(db, stoklar)

    for stok in stoklar:
        toplam = len(stok.seriler)
        sayilan_seriler = [s for s in stok.seriler if s.sayildi]
        sayilan = len(sayilan_seriler)
        fark = sayilan - stok.portal_sayim

        baslik_renk = ACCENT_BG if stok.sonradan_eklendi else GRUP_BASLIK_BG
        for col in range(1, len(kolonlar) + 1):
            c = ws.cell(row=satir, column=col)
            c.fill = PatternFill("solid", fgColor=baslik_renk)
            c.border = _border()
            c.font = Font(bold=True, size=11)

        ws.cell(row=satir, column=1, value=stok.stok_kodu)
        ws.cell(row=satir, column=2, value=stok.urun_adi)
        ws.cell(row=satir, column=3, value=f"Adet: {toplam}")
        ws.cell(row=satir, column=4, value=stok.portal_sayim)
        ws.cell(row=satir, column=5, value=sayilan)
        farkc = ws.cell(row=satir, column=6, value=fark)
        if fark == 0:
            farkc.fill = PatternFill("solid", fgColor=GOOD_BG); farkc.font = Font(bold=True, color="FF2E7D32")
        elif fark > 0:
            farkc.fill = PatternFill("solid", fgColor=WARN_BG); farkc.font = Font(bold=True, color="FFA9521A")
        else:
            farkc.fill = PatternFill("solid", fgColor=BAD_BG); farkc.font = Font(bold=True, color="FF8A2222")

        if stok.sonradan_eklendi:
            ws.cell(row=satir, column=10, value="EVET")
        satir += 1

        for seri in sorted(stok.seriler, key=lambda x: x.seri_no):
            sayan_ad = adlar.get(seri.sayan_id) if seri.sayan_id else None
            zimmet_ad = _zimmet_ad(seri, adlar, employee_adlar)

            durum = "Sayilmadi"
            durum_bg = None
            if seri.cikis_zaman:
                durum = "Cikis"; durum_bg = BAD_BG
            elif _zimmetli(seri):
                durum = "Zimmet"; durum_bg = WARN_BG
            elif seri.sayildi:
                durum = "Sayildi"; durum_bg = GOOD_BG

            ws.cell(row=satir, column=3, value=seri.seri_no)
            ws.cell(row=satir, column=7,
                    value=seri.sayim_tarihi.strftime("%d.%m.%Y %H:%M:%S") if seri.sayim_tarihi else "")
            ws.cell(row=satir, column=8, value=sayan_ad or "")
            dc = ws.cell(row=satir, column=9, value=durum)
            if durum_bg:
                dc.fill = PatternFill("solid", fgColor=durum_bg)
                dc.font = Font(bold=True, size=10)
            if seri.sonradan_eklendi:
                ec = ws.cell(row=satir, column=10, value="EVET")
                ec.fill = PatternFill("solid", fgColor=ACCENT_BG)
                ec.font = Font(bold=True, color=ACCENT)
            ws.cell(row=satir, column=11, value=seri.cikis_notu or "")
            ws.cell(row=satir, column=12, value=zimmet_ad or "")
            ws.cell(row=satir, column=13, value=seri.zimmet_notu or "")
            ws.cell(row=satir, column=14, value=seri.notlar or "")

            for col in range(1, len(kolonlar) + 1):
                cc = ws.cell(row=satir, column=col)
                cc.border = _border()
                if col == 3 and seri.sonradan_eklendi:
                    cc.font = Font(bold=True, italic=True, color=ACCENT)
            detay_satiri_grupla(ws, satir)
            satir += 1


def _ozet_sayfasi(db: Session, wb: Workbook, oturum: SayimOturumu) -> None:
    ws = wb.create_sheet("Ozet", 0)

    kolonlar = [
        ("Stok Kodu", 14), ("Urun Adi", 36), ("Portal Sayi", 12),
        ("Sayilan", 12), ("Fark", 10), ("Yuzde", 10),
        ("Sonradan Eklendi", 13), ("Cikis", 8), ("Zimmet", 8),
    ]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(kolonlar))
    c0 = ws.cell(row=1, column=1, value=f"OZET — {oturum.ad}")
    c0.font = Font(bold=True, color="FFFFFFFF", size=14)
    c0.fill = PatternFill("solid", fgColor=DEEP)
    c0.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    stoklar = db.execute(
        select(DepoSayimStok).where(DepoSayimStok.oturum_id == oturum.id).order_by(DepoSayimStok.stok_kodu)
        .options(selectinload(DepoSayimStok.seriler))
    ).scalars().all()
    toplam_portal = sum(s.portal_sayim for s in stoklar)
    toplam_sayilan = sum(1 for s in stoklar for x in s.seriler if x.sayildi)
    toplam_seri = sum(len(s.seriler) for s in stoklar)
    toplam_cikis = sum(1 for s in stoklar for x in s.seriler if x.cikis_zaman)
    toplam_zimmet = sum(1 for s in stoklar for x in s.seriler if _zimmetli(x))
    sonradan_seri = sum(1 for s in stoklar for x in s.seriler if x.sonradan_eklendi)
    sonradan_stok = sum(1 for s in stoklar if s.sonradan_eklendi)

    metrikler = [
        ("Stok Kalemi", len(stoklar), DEEP),
        ("Toplam Seri", toplam_seri, DEEP),
        ("Sayilan", toplam_sayilan, GOOD),
        ("Kalan", toplam_seri - toplam_sayilan, ACCENT),
        ("Portalda Toplam", toplam_portal, DEEP),
        ("Portal Fark", toplam_sayilan - toplam_portal, ACCENT),
        ("Sonradan Eklenen Stok", sonradan_stok, ACCENT),
        ("Sonradan Eklenen Seri", sonradan_seri, ACCENT),
        ("Cikis", toplam_cikis, BAD),
        ("Zimmet", toplam_zimmet, ACCENT),
    ]

    for i, (etiket, deger, renk) in enumerate(metrikler):
        col = (i % 5) * 2 + 1
        sat = 3 + (i // 5) * 3
        et = ws.cell(row=sat, column=col, value=etiket)
        et.font = Font(bold=True, size=9, color="FF777777")
        et.alignment = Alignment(horizontal="left")
        dg = ws.cell(row=sat + 1, column=col, value=deger)
        dg.font = Font(bold=True, size=18, color=renk)
        dg.alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=sat, start_column=col, end_row=sat, end_column=col + 1)
        ws.merge_cells(start_row=sat + 1, start_column=col, end_row=sat + 1, end_column=col + 1)

    bas_satir = 10
    _baslik_yaz(ws, bas_satir, kolonlar)
    _kolon_genislikleri(ws, kolonlar)
    ws.row_dimensions[bas_satir].height = 28
    ws.freeze_panes = f"A{bas_satir + 1}"

    sat = bas_satir + 1
    sirali = sorted(
        stoklar,
        key=lambda s: abs(sum(1 for x in s.seriler if x.sayildi) - s.portal_sayim),
        reverse=True,
    )
    for stok in sirali:
        sayilan = sum(1 for x in stok.seriler if x.sayildi)
        cikis = sum(1 for x in stok.seriler if x.cikis_zaman)
        zimmet = sum(1 for x in stok.seriler if _zimmetli(x))
        fark = sayilan - stok.portal_sayim
        yuzde = round(sayilan / stok.portal_sayim * 100, 1) if stok.portal_sayim > 0 else 0

        ws.cell(row=sat, column=1, value=stok.stok_kodu)
        ws.cell(row=sat, column=2, value=stok.urun_adi)
        ws.cell(row=sat, column=3, value=stok.portal_sayim)
        ws.cell(row=sat, column=4, value=sayilan)
        farkc = ws.cell(row=sat, column=5, value=fark)
        ws.cell(row=sat, column=6, value=f"{yuzde}%")
        if stok.sonradan_eklendi:
            ec = ws.cell(row=sat, column=7, value="EVET")
            ec.fill = PatternFill("solid", fgColor=ACCENT_BG)
            ec.font = Font(bold=True, color=ACCENT)
        ws.cell(row=sat, column=8, value=cikis or "")
        ws.cell(row=sat, column=9, value=zimmet or "")

        if fark == 0:
            farkc.fill = PatternFill("solid", fgColor=GOOD_BG); farkc.font = Font(bold=True, color="FF2E7D32")
        elif fark > 0:
            farkc.fill = PatternFill("solid", fgColor=WARN_BG); farkc.font = Font(bold=True, color="FFA9521A")
        else:
            farkc.fill = PatternFill("solid", fgColor=BAD_BG); farkc.font = Font(bold=True, color="FF8A2222")

        for col in range(1, len(kolonlar) + 1):
            ws.cell(row=sat, column=col).border = _border()
        sat += 1


def _log_sayfasi(db: Session, wb: Workbook, oturum: SayimOturumu) -> None:
    from app.models import DepoTaramaLog
    ws = wb.create_sheet("Tarama Log")

    kolonlar = [
        ("Zaman", 19), ("Kullanici", 16), ("Seri Giris", 22),
        ("Durum", 12), ("Stok Kodu", 14), ("Urun Adi", 32), ("Aciklama", 30),
    ]
    _baslik_yaz(ws, 1, kolonlar)
    _kolon_genislikleri(ws, kolonlar)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    rows = db.execute(
        select(DepoTaramaLog, DepoUser.ad)
        .outerjoin(DepoUser, DepoUser.id == DepoTaramaLog.kullanici_id)
        .where(DepoTaramaLog.oturum_id == oturum.id)
        .order_by(DepoTaramaLog.zaman.asc())
    ).all()
    durum_renk = {
        "basarili": GOOD_BG, "mukerrer": WARN_BG, "bulunamadi": BAD_BG, "cakisma": WARN_BG,
    }
    sat = 2
    for log, ad in rows:
        ws.cell(row=sat, column=1, value=log.zaman.strftime("%d.%m.%Y %H:%M:%S"))
        ws.cell(row=sat, column=2, value=ad or "")
        ws.cell(row=sat, column=3, value=log.seri_giris)
        dc = ws.cell(row=sat, column=4, value=log.durum)
        bg = durum_renk.get(log.durum)
        if bg:
            dc.fill = PatternFill("solid", fgColor=bg)
            dc.font = Font(bold=True)
        ws.cell(row=sat, column=5, value=log.stok_kodu or "")
        ws.cell(row=sat, column=6, value=log.urun_adi or "")
        ws.cell(row=sat, column=7, value=log.aciklama or "")
        for col in range(1, len(kolonlar) + 1):
            ws.cell(row=sat, column=col).border = _border()
        sat += 1
